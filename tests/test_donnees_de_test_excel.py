"""Lot PT-02a (A37) — le tableau des données de test en Excel : exporter, et importer en COMPLÉTANT.

Ce que ces tests gardent :
1. l'import complète : variables et jeux ajoutés, cases remplies mises à jour ; une case vide ne vide rien, rien n'est supprimé ;
2. une seule erreur dans le fichier (nommée par sa case) : RIEN n'est écrit ;
3. les noms et les valeurs suivent les règles du remplacement (un fichier ne fait pas entrer ce qu'un test ne saurait recevoir) ;
4. l'export se relit à l'identique, et une valeur « =… » y reste du texte, jamais une formule.
"""

from __future__ import annotations

import datetime as dt
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from testpilot import config
from testpilot.api import app as app_mod
from testpilot.api.services import donnees_de_test_excel as excel
from testpilot.store.db import get_initialized_db
from testpilot.store.repositories import DonneeDeTestInvalide, ProjectRepo, TestDataRepo


@pytest.fixture
def base(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    conn = get_initialized_db(tmp_path / "db.sqlite")
    yield conn
    conn.close()


def _classeur(*lignes) -> bytes:
    classeur = openpyxl.Workbook()
    for ligne in lignes:
        classeur.active.append(list(ligne))
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def _projet(conn) -> int:
    """Le tableau de la question posée au porteur : client = Martin, ville = Lyon (Par défaut)."""
    pid = ProjectRepo(conn).create(name="Boutique", connector_type="web", base_url="http://app")
    depot = TestDataRepo(conn)
    ids = {nom: depot.creer_variable(pid, nom=nom) for nom in ("client", "ville")}
    depot.poser_valeurs(pid, depot.jeux(pid)[0]["id"], {ids["client"]: "Martin", ids["ville"]: "Lyon"})
    return pid


def _tableau(conn, pid) -> dict[tuple[str, str], str]:
    t = TestDataRepo(conn).tableau(pid)
    noms = {v["id"]: v["nom"] for v in t["variables"]}
    return {(noms[vid], j["nom"]): valeur for j in t["jeux"] for vid, valeur in t["valeurs"][j["id"]].items()}


def test_falsifiable_l_import_complete_le_tableau_sans_rien_vider(base):
    pid = _projet(base)
    fichier = _classeur(("Variable", "Par défaut", "Grosse commande"),
                        ("client", "Durand", None),
                        ("quantite", None, 12),
                        ("ville", None, None))

    changements = excel.appliquer(base, pid, excel.lire(fichier))

    assert changements["nouvelles_variables"] == ["quantite"] and changements["nouveaux_jeux"] == ["Grosse commande"]
    assert _tableau(base, pid) == {("client", "Par défaut"): "Durand", ("ville", "Par défaut"): "Lyon",
                                   ("quantite", "Grosse commande"): "12"}


def test_l_apercu_annonce_exactement_les_changements_sans_rien_ecrire(base):
    pid = _projet(base)
    fichier = _classeur(("Variable", "par défaut"), ("client", "Durand"), ("ville", "Lyon"))

    vu = excel.comparer(base, pid, excel.lire(fichier))

    assert vu == {"nouvelles_variables": [], "nouveaux_jeux": [], "erreurs": [],
                  "cases": [{"variable": "client", "jeu": "Par défaut", "avant": "Martin", "apres": "Durand"}]}
    assert _tableau(base, pid)[("client", "Par défaut")] == "Martin", "l'aperçu n'écrit rien"


@pytest.mark.parametrize("lignes, erreur", [
    ((("Nom", "Par défaut"), ("client", "Durand")), "case A1"),
    ((("Variable", "Par défaut"), ("Client Final", "Durand")), "case A2"),
    ((("Variable", "Par défaut"), ("client", "Durand {{ville}}")), "case B2"),
    ((("Variable", "Par défaut"), ("client", "<nom>")), "case B2"),
    ((("Variable", "Export", "export"), ("client", "Durand")), "case C1"),
    ((("Variable", "Par défaut"), ("client", "Durand"), ("client", "Dupont")), "case A3"),
    ((("Variable", "Par défaut"), ("client", "Durand", "orphelin")), "case C2"),
    ((("Variable", "Par défaut", None, "Export"), ("client", "Durand")), "case C1"),
])
def test_falsifiable_une_seule_erreur_et_rien_n_est_ecrit(base, lignes, erreur):
    pid = _projet(base)
    lecture = excel.lire(_classeur(*lignes))

    assert any(e.startswith(erreur) for e in lecture.erreurs), lecture.erreurs
    with pytest.raises(DonneeDeTestInvalide):
        excel.appliquer(base, pid, lecture)
    assert _tableau(base, pid) == {("client", "Par défaut"): "Martin", ("ville", "Par défaut"): "Lyon"}
    assert [j["nom"] for j in TestDataRepo(base).jeux(pid)] == ["Par défaut"]


def test_falsifiable_les_plafonds_du_projet_sont_tenus(base, monkeypatch):
    pid = _projet(base)
    monkeypatch.setattr(excel, "_JEUX_MAX", 2)
    fichier = _classeur(("Variable", "Par défaut", "Export", "Grosse commande"), ("client", "Durand", "A", "B"))

    with pytest.raises(DonneeDeTestInvalide, match="jeux de données"):
        excel.appliquer(base, pid, excel.lire(fichier))
    assert len(TestDataRepo(base).jeux(pid)) == 1


def test_les_cases_non_texte_sont_ecrites_comme_on_les_lit(base):
    pid = _projet(base)
    fichier = _classeur(("Variable", "Par défaut"), ("client", 12.5), ("ville", dt.datetime(2026, 10, 14)),
                        ("quantite", 3.0))

    excel.appliquer(base, pid, excel.lire(fichier))

    t = _tableau(base, pid)
    assert (t[("client", "Par défaut")], t[("ville", "Par défaut")], t[("quantite", "Par défaut")]) == ("12.5", "2026-10-14", "3")


def test_un_projet_sans_variable_recoit_le_tableau_entier(base):
    pid = ProjectRepo(base).create(name="Neuf", connector_type="web", base_url="http://app")

    excel.appliquer(base, pid, excel.lire(_classeur(("Variable", "Par défaut", "Export"), ("client", "Martin", "Durand"))))

    assert [j["nom"] for j in TestDataRepo(base).jeux(pid)] == ["Par défaut", "Export"]
    assert _tableau(base, pid) == {("client", "Par défaut"): "Martin", ("client", "Export"): "Durand"}


def test_l_export_se_relit_a_l_identique(base):
    pid = _projet(base)
    depot = TestDataRepo(base)
    export = depot.creer_jeu(pid, nom="Export")
    depot.poser_valeurs(pid, export, {depot.variables(pid)[0]["id"]: "Durand"})

    vu = excel.comparer(base, pid, excel.lire(excel.exporter(base, pid)))

    assert vu == {"nouvelles_variables": [], "nouveaux_jeux": [], "cases": [], "erreurs": []}


def test_falsifiable_une_valeur_qui_commence_par_egal_reste_du_texte_a_l_export(base):
    pid = _projet(base)
    depot = TestDataRepo(base)
    depot.poser_valeurs(pid, depot.jeux(pid)[0]["id"], {depot.variables(pid)[0]["id"]: "=1+1"})

    case = openpyxl.load_workbook(io.BytesIO(excel.exporter(base, pid))).active["B2"]

    assert (case.value, case.data_type) == ("=1+1", "s")


def test_un_fichier_qui_n_est_pas_un_classeur_est_refuse(base):
    with pytest.raises(excel.FichierExcelInvalide):
        excel.lire(b"pas un classeur")


# ── Par l'API ────────────────────────────────────────────────────────────────────────────────────────────────────────────


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "api.db")
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    return TestClient(app_mod.app)


def test_apercu_puis_import_par_l_api(client):
    pid = client.post("/api/projects", json={"name": "B", "connector_type": "web", "base_url": "https://a.test",
                                              "username": "u", "password": "p"}).json()["id"]
    base = f"/api/projects/{pid}/donnees-de-test"
    fichier = {"file": ("donnees.xlsx", _classeur(("Variable", "Par défaut"), ("client", "Durand")),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    vu = client.post(f"{base}/import/apercu", files=fichier)
    assert vu.status_code == 200, vu.text
    assert vu.json()["nouvelles_variables"] == ["client"] and client.get(base).json()["variables"] == []

    fait = client.post(f"{base}/import", files=fichier)
    assert fait.status_code == 200, fait.text
    assert [v["nom"] for v in fait.json()["variables"]] == ["client"]

    export = client.get(f"{base}/export")
    assert export.status_code == 200 and export.headers["content-type"].startswith("application/vnd.openxmlformats")


def test_falsifiable_un_import_en_erreur_est_refuse_par_l_api(client):
    pid = client.post("/api/projects", json={"name": "B", "connector_type": "web", "base_url": "https://a.test",
                                              "username": "u", "password": "p"}).json()["id"]
    base = f"/api/projects/{pid}/donnees-de-test"
    fichier = {"file": ("d.xlsx", _classeur(("Variable", "Par défaut"), ("client", "{{x}}")), "application/octet-stream")}

    r = client.post(f"{base}/import", files=fichier)

    assert r.status_code == 422 and "case B2" in r.text
    assert client.get(base).json()["variables"] == []


# ── Revue du 2026-10-14 ───────────────────────────────────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("vide", ["   ", "\xa0", "\t"])
def test_falsifiable_une_case_d_espaces_ne_vide_rien(base, vide):
    pid = _projet(base)

    excel.appliquer(base, pid, excel.lire(_classeur(("Variable", "Par défaut"), ("client", vide))))

    assert _tableau(base, pid)[("client", "Par défaut")] == "Martin"


def test_falsifiable_le_plafond_de_jeux_du_projet_compte_ceux_qui_existent(base, monkeypatch):
    """Le fichier seul reste sous le plafond ; c'est l'ajout au projet qui le dépasse."""
    pid = _projet(base)
    TestDataRepo(base).creer_jeu(pid, nom="Export")
    monkeypatch.setattr(excel, "_JEUX_MAX", 3)
    lecture = excel.lire(_classeur(("Variable", "Grosse commande", "Retour"), ("client", "A", "B")))
    assert lecture.erreurs == []

    with pytest.raises(DonneeDeTestInvalide, match="jeux de données"):
        excel.appliquer(base, pid, lecture)
    assert len(TestDataRepo(base).jeux(pid)) == 2


def test_falsifiable_le_plafond_de_variables_du_projet_compte_celles_qui_existent(base, monkeypatch):
    pid = _projet(base)   # client, ville
    monkeypatch.setattr(excel, "_VARIABLES_MAX", 3)
    lecture = excel.lire(_classeur(("Variable", "Par défaut"), ("pays", "FR"), ("quantite", "2")))
    assert lecture.erreurs == []

    with pytest.raises(DonneeDeTestInvalide, match="variables"):
        excel.appliquer(base, pid, lecture)
    assert len(TestDataRepo(base).variables(pid)) == 2


def test_une_ligne_vide_mise_en_forme_ne_compte_pas(base):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(["Variable", "Par défaut"])
    feuille.append(["client", "Durand"])
    feuille.row_dimensions[500].height = 30
    feuille.cell(row=500, column=3).number_format = "0.00"
    tampon = io.BytesIO()
    classeur.save(tampon)

    lecture = excel.lire(tampon.getvalue())

    assert lecture.erreurs == [] and lecture.variables == ["client"]


@pytest.mark.parametrize("valeur, erreur", [("#N/A", "en erreur"), (12345678901234567890, "nombre trop long")])
def test_falsifiable_une_case_en_erreur_ou_un_nombre_trop_long_est_refuse(base, valeur, erreur):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(["Variable", "Par défaut"])
    feuille.append(["client", valeur])
    if valeur == "#N/A":
        feuille["B2"].data_type = "e"
    tampon = io.BytesIO()
    classeur.save(tampon)

    lecture = excel.lire(tampon.getvalue())

    assert any(e.startswith("case B2") and erreur in e for e in lecture.erreurs), lecture.erreurs


def test_falsifiable_par_defaut_ecrit_en_deux_caracteres_ne_cree_pas_un_second_jeu(base):
    import unicodedata

    pid = _projet(base)
    nfd = unicodedata.normalize("NFD", "Par défaut")
    assert nfd != "Par défaut"

    excel.appliquer(base, pid, excel.lire(_classeur(("Variable", nfd), ("client", "Durand"))))

    assert [j["nom"] for j in TestDataRepo(base).jeux(pid)] == ["Par défaut"]
    assert _tableau(base, pid)[("client", "Par défaut")] == "Durand"


def test_une_feuille_corrompue_est_un_fichier_illisible_jamais_une_erreur_serveur(base):
    import re
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(_classeur(("Variable", "Par défaut"), ("client", 5))))
    sortie = io.BytesIO()
    with zipfile.ZipFile(sortie, "w") as cible:
        for info in source.infolist():
            contenu = source.read(info.filename)
            if info.filename.endswith("sheet1.xml"):
                contenu = re.sub(rb"<v>5</v>", b"<v>abc</v>", contenu)
            cible.writestr(info, contenu)

    with pytest.raises(excel.FichierExcelInvalide):
        excel.lire(sortie.getvalue())


def test_falsifiable_une_archive_qui_gonfle_est_refusee_avant_d_etre_ouverte(base, monkeypatch):
    monkeypatch.setattr(excel, "DECOMPRESSE_MAX", 1000)

    with pytest.raises(excel.FichierExcelInvalide, match="trop volumineux"):
        excel.lire(_classeur(("Variable", "Par défaut"), ("client", "x" * 5000)))


def test_le_modele_exporte_garde_le_texte(base):
    pid = _projet(base)

    feuille = openpyxl.load_workbook(io.BytesIO(excel.exporter(base, pid))).active

    assert feuille["B2"].number_format == "@"
